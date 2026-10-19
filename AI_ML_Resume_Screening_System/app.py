import os
import re
import io
import csv
import json
import tempfile
from datetime import datetime
from flask import Flask, render_template, request, jsonify, send_file
from werkzeug.utils import secure_filename
from pypdf import PdfReader
from docx import Document
import numpy as np
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from sklearn.pipeline import Pipeline
from sklearn.linear_model import LogisticRegression
from sklearn.model_selection import train_test_split
from sklearn.metrics import accuracy_score
import joblib

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 12 * 1024 * 1024
ALLOWED_EXTENSIONS = {"pdf", "docx", "txt"}
MODEL_PATH = os.path.join(os.path.dirname(__file__), "models", "demo_suitability_model.joblib")
os.makedirs(os.path.dirname(MODEL_PATH), exist_ok=True)

# This small illustrative dataset is for demonstrating the ML pipeline only.
# Do NOT use a model trained on synthetic examples to make real hiring decisions.
TRAINING_EXAMPLES = [
    ("python sql pandas machine learning data analysis dashboard", "python sql pandas data analyst reporting", 1),
    ("built REST API with flask python unit tests", "python flask backend API development", 1),
    ("trained classification model with scikit learn feature engineering", "machine learning scikit-learn model training", 1),
    ("managed customer calls sales targets communication", "deep learning pytorch computer vision", 0),
    ("graphic design branding photoshop illustrator", "python SQL data analytics", 0),
    ("retail cashier inventory customer service", "NLP transformers Python AI engineer", 0),
    ("react typescript css responsive web application", "Java Spring Boot microservices", 0),
    ("SQL queries pandas cleaning data visualization Power BI", "SQL pandas data analysis visualization", 1),
    ("computer vision YOLO object detection OpenCV", "YOLO OpenCV computer vision", 1),
    ("Java OOP algorithms data structures", "Python NLP sentence transformers", 0),
    ("Excel pivot tables business reporting", "Excel reporting business intelligence", 1),
    ("content writing social media campaigns", "cloud infrastructure kubernetes devops", 0),
    ("natural language processing embeddings semantic search", "NLP embeddings semantic search", 1),
    ("network troubleshooting hardware support", "predictive modeling machine learning", 0),
    ("built dashboard using Tableau SQL and Python", "Tableau SQL Python dashboard", 1),
    ("C++ embedded firmware microcontrollers", "marketing analytics Python pandas", 0),
]

def build_demo_model():
    texts = [f"RESUME: {r} JOB: {j}" for r, j, _ in TRAINING_EXAMPLES]
    labels = [label for _, _, label in TRAINING_EXAMPLES]
    pipe = Pipeline([
        ("tfidf", TfidfVectorizer(ngram_range=(1, 2), max_features=2500, sublinear_tf=True)),
        ("classifier", LogisticRegression(max_iter=1000, class_weight="balanced", random_state=42))
    ])
    pipe.fit(texts, labels)
    joblib.dump(pipe, MODEL_PATH)
    return pipe

def get_ml_model():
    if os.path.exists(MODEL_PATH):
        try:
            return joblib.load(MODEL_PATH)
        except Exception:
            pass
    return build_demo_model()

ML_MODEL = get_ml_model()

def extract_text(filename, data):
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "txt":
        return data.decode("utf-8", errors="ignore")
    if ext == "pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if ext == "docx":
        doc = Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(" ".join(cell.text for cell in row.cells))
        return "\n".join(paragraphs)
    raise ValueError("Unsupported file type")

def clean_text(text):
    return re.sub(r"\s+", " ", (text or "")).strip()

def extract_skills(text):
    skills = [
        "python", "java", "c++", "sql", "pandas", "numpy", "scikit-learn",
        "machine learning", "deep learning", "nlp", "natural language processing",
        "sentence transformers", "tensorflow", "pytorch", "flask", "django",
        "html", "css", "javascript", "react", "typescript", "power bi", "tableau",
        "excel", "data analysis", "data visualization", "statistics", "git",
        "github", "docker", "kubernetes", "aws", "azure", "computer vision",
        "opencv", "yolo", "rest api", "fastapi", "linux", "data structures",
        "algorithms", "communication", "project management", "generative ai",
        "large language models", "rag", "vector databases", "mongodb", "postgresql"
    ]
    lower = text.lower()
    found = []
    for skill in skills:
        if skill in lower:
            found.append(skill)
    return sorted(set(found))

def semantic_similarity(resume_text, job_text):
    """Use Sentence Transformers if installed and model download/cache is available.
    Otherwise use a local TF-IDF similarity fallback so the app remains runnable.
    """
    try:
        from sentence_transformers import SentenceTransformer
        global _embedding_model
        try:
            _embedding_model
        except NameError:
            _embedding_model = SentenceTransformer("sentence-transformers/all-MiniLM-L6-v2")
        vectors = _embedding_model.encode([resume_text, job_text], normalize_embeddings=True)
        score = float(np.dot(vectors[0], vectors[1]))
        return max(0.0, min(1.0, score)), "Sentence Transformers (all-MiniLM-L6-v2)"
    except Exception:
        vectorizer = TfidfVectorizer(stop_words="english", ngram_range=(1, 2), max_features=12000)
        matrix = vectorizer.fit_transform([resume_text, job_text])
        score = float(cosine_similarity(matrix[0:1], matrix[1:2])[0][0])
        return score, "TF-IDF fallback (Sentence Transformers unavailable)"

def analyze_resume(filename, resume_text, job_text, required_skills):
    resume_text = clean_text(resume_text)
    job_text = clean_text(job_text)
    if not resume_text:
        raise ValueError(f"Could not extract readable text from {filename}.")
    resume_skills = extract_skills(resume_text)
    job_skills = extract_skills(job_text)
    if required_skills:
        target_skills = required_skills
    else:
        target_skills = job_skills
    matched = sorted(set(resume_skills).intersection(target_skills))
    missing = sorted(set(target_skills).difference(resume_skills))
    skill_score = (len(matched) / max(1, len(target_skills))) if target_skills else 0.0
    semantic_score, semantic_method = semantic_similarity(resume_text[:20000], job_text[:10000])
    ml_input = f"RESUME: {resume_text[:12000]} JOB: {job_text[:6000]}"
    ml_probability = float(ML_MODEL.predict_proba([ml_input])[0][1])
    # Transparent blended score: semantic relevance, skill coverage, and demo classifier probability.
    overall = round(100 * (0.50 * semantic_score + 0.30 * skill_score + 0.20 * ml_probability))
    return {
        "filename": filename,
        "score": max(0, min(100, overall)),
        "semantic_score": round(semantic_score * 100),
        "skill_score": round(skill_score * 100),
        "ml_probability": round(ml_probability * 100),
        "matched_skills": matched,
        "missing_skills": missing,
        "resume_skills": resume_skills,
        "target_skills": target_skills,
        "semantic_method": semantic_method,
        "preview": resume_text[:360] + ("…" if len(resume_text) > 360 else "")
    }

@app.route("/")
def index():
    return render_template("index.html")

@app.route("/api/analyze", methods=["POST"])
def api_analyze():
    job_description = clean_text(request.form.get("job_description", ""))
    required_text = clean_text(request.form.get("required_skills", ""))
    if len(job_description) < 20:
        return jsonify({"error": "Please enter a job description of at least 20 characters."}), 400
    uploads = request.files.getlist("resumes")
    uploads = [f for f in uploads if f and f.filename]
    if not uploads:
        return jsonify({"error": "Upload at least one resume file."}), 400
    if len(uploads) > 10:
        return jsonify({"error": "Please upload no more than 10 resumes at once."}), 400
    required_skills = sorted(set(s.strip().lower() for s in re.split(r"[,;\n]", required_text) if s.strip()))
    results, errors = [], []
    for file in uploads:
        filename = secure_filename(file.filename) or "resume"
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        if ext not in ALLOWED_EXTENSIONS:
            errors.append(f"{filename}: unsupported format.")
            continue
        try:
            data = file.read()
            text = extract_text(filename, data)
            results.append(analyze_resume(filename, text, job_description, required_skills))
        except Exception as exc:
            errors.append(f"{filename}: {str(exc)}")
    results.sort(key=lambda x: x["score"], reverse=True)
    return jsonify({
        "results": results,
        "errors": errors,
        "count": len(results),
        "average_score": round(sum(r["score"] for r in results) / len(results)) if results else 0,
        "model_note": "The supervised ML classifier is trained on small synthetic examples for demonstration only; it is not validated for hiring."
    })

@app.route("/api/demo", methods=["GET"])
def api_demo():
    return jsonify({
        "job_description": "We are looking for a junior AI/Data Analyst with Python, SQL, pandas, data visualization, statistics, and machine learning fundamentals. The candidate should be able to clean data, build dashboards, evaluate models, communicate findings, and work with Git.",
        "required_skills": "Python, SQL, pandas, data visualization, statistics, machine learning, Git"
    })

@app.route("/api/export", methods=["POST"])
def api_export():
    payload = request.get_json(silent=True) or {}
    results = payload.get("results", [])
    if not results:
        return jsonify({"error": "No results available to export."}), 400
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Rank", "Resume", "Overall Score", "Semantic Similarity", "Skill Coverage", "Demo ML Probability", "Matched Skills", "Missing Skills"])
    for rank, row in enumerate(results, start=1):
        writer.writerow([rank, row.get("filename", ""), row.get("score", ""), row.get("semantic_score", ""), row.get("skill_score", ""), row.get("ml_probability", ""), ", ".join(row.get("matched_skills", [])), ", ".join(row.get("missing_skills", []))])
    mem = io.BytesIO(output.getvalue().encode("utf-8-sig"))
    mem.seek(0)
    return send_file(mem, mimetype="text/csv", as_attachment=True, download_name=f"resume_screening_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv")

if __name__ == "__main__":
    print("Starting AI + ML Resume Screening System at http://127.0.0.1:5000")
    app.run(debug=True)
